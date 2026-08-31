import asyncio
from datetime import datetime
from io import BytesIO
from typing import List, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from fastapi_pagination import Page
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session, joinedload

from core.config import is_survey_export_column_enabled
from core.time import as_utc, resolve_timezone, utc_isoformat, utc_now
from features.feedback.filtering import FilterRequest, build_survey_query, get_filter_params
from features.identity.service.security import get_current_actor, require_admin
from infrastructure.database.dbo.Channel import Channel
from infrastructure.database.dbo.DeliveryService import DeliveryService
from infrastructure.database.dbo.Department import Department
from infrastructure.database.dbo.Keyword import Keyword
from infrastructure.database.dbo.Store import Store
from infrastructure.database.dbo.Survey import Survey
from infrastructure.database.dbo.SurveyDepartments import SurveyDepartments
from infrastructure.database.dbo.SurveyKeywords import SurveyKeywords
from infrastructure.database.dbo.SurveyTopics import SurveyTopics
from infrastructure.database.dbo.Topic import Topic
from infrastructure.database.session import get_db

router = APIRouter(
    prefix="/surveys",
    tags=["surveys"],
    dependencies=[Depends(get_db), Depends(get_current_actor)],
)


class StoreResponse(BaseModel):
    store_key: int
    store_name_english: Optional[str] = None
    store_name_local: Optional[str] = None
    bu_key: Optional[str] = None
    area_manager: Optional[str] = None
    store_format: Optional[str] = None
    store_type: Optional[str] = None
    operations_controller: Optional[str] = None
    regional_manager: Optional[str] = None
    px: Optional[str] = None
    csr: Optional[str] = None
    dr: Optional[str] = None
    mag_type: Optional[str] = None
    cf_grouping: Optional[str] = None
    store_brand: Optional[str] = None
    competitor: Optional[str] = None
    region: Optional[str] = None
    area: Optional[str] = None
    province: Optional[str] = None
    territory: Optional[str] = None
    toh: Optional[str] = None
    district: Optional[str] = None
    city: Optional[str] = None
    operations_manager: Optional[str] = None
    district_manager: Optional[str] = None
    sic: Optional[str] = None
    soc: Optional[str] = None
    tech_life_type: Optional[str] = None
    operation_manager_tl: Optional[str] = None
    region_manager_tl: Optional[str] = None
    relocation: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    store_open_date: Optional[str] = None
    store_close_date: Optional[str] = None
    is_closed: bool


class ChannelResponse(BaseModel):
    id: int
    name: str


class DeliveryServiceResponse(BaseModel):
    id: int
    name: str


class DepartmentWithSentimentResponse(BaseModel):
    department_id: int
    name: str
    sentiment: str


class KeywordWithSentimentResponse(BaseModel):
    keyword_id: int
    keyword: str
    sentiment: str


class TopicWithSentimentResponse(BaseModel):
    topic_id: int
    topic: str
    sentiment: str


class SurveyResponse(BaseModel):
    id: int
    store: StoreResponse
    channel: Optional[ChannelResponse]
    delivery_service: Optional[DeliveryServiceResponse]
    departments: List[DepartmentWithSentimentResponse]
    topics: List[TopicWithSentimentResponse]
    keywords: List[KeywordWithSentimentResponse]
    comment: str
    sentiment: str
    topic_sentiment: str
    topic_sentiment_score: float
    cls: Optional[float] = None
    reported_at: str
    created_at: str
    updated_at: str


@router.get("")
async def get_surveys(
    filter_params: FilterRequest = Depends(get_filter_params),
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=100),
    db: Session = Depends(get_db),
) -> Page[SurveyResponse]:
    filter_dict = filter_params.model_dump()
    
    # First get distinct survey IDs that match the filters
    # Include reported_at in select for ORDER BY compatibility with DISTINCT
    id_query = build_survey_query(db.query(Survey.id, Survey.reported_at).distinct(), filter_dict)
    id_query = id_query.order_by(Survey.reported_at.desc())
    
    # Calculate total count
    total = id_query.count()
    
    # Apply pagination to get survey IDs
    offset = (page - 1) * size
    survey_ids = [row[0] for row in id_query.offset(offset).limit(size).all()]
    
    # Now fetch full survey objects for these IDs with proper eager loading
    surveys = (
        db.query(Survey)
        .filter(Survey.id.in_(survey_ids))
        .options(
            joinedload(Survey.store),
            joinedload(Survey.survey_topics),
            joinedload(Survey.survey_keywords),
            joinedload(Survey.survey_departments),
            joinedload(Survey.channel),
            joinedload(Survey.delivery_service)
        )
        .order_by(Survey.reported_at.desc())
        .all()
    ) if survey_ids else []

    # Convert to response models
    survey_responses = [
        SurveyResponse.model_validate(
            survey.to_dict(filter_params.timezone)
        )
        for survey in surveys
    ]

    # Create pagination response
    from fastapi_pagination import Params

    params = Params(page=page, size=size)

    return Page.create(items=survey_responses, total=total, params=params)


class CreateSurveyDepartmentRequest(BaseModel):
    name: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyTopicRequest(BaseModel):
    topic: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyKeywordRequest(BaseModel):
    keyword: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyRequest(BaseModel):
    store_key: int
    channel: Optional[str] = None
    delivery_service: Optional[str] = None
    departments: List[CreateSurveyDepartmentRequest]
    comment: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"
    topics: List[CreateSurveyTopicRequest]
    keywords: List[CreateSurveyKeywordRequest]
    reported_at: datetime = Field(default_factory=utc_now)
    cls: Optional[float] = None


@router.post("/")
async def create_survey(
    survey_request: CreateSurveyRequest,
    db: Session = Depends(get_db),
    _: object = Depends(require_admin),
):

    # Check if store exists
    store = db.query(Store).filter(Store.store_key == survey_request.store_key).first()
    if not store:
        raise HTTPException(status_code=404, detail="Store not found")
    # Check if departments exist
    departments = (
        db.query(Department)
        .filter(Department.name.in_([dept.name for dept in survey_request.departments]))
        .all()
    )
    if not departments:
        raise HTTPException(status_code=404, detail="Departments not found")
    # Check if topics exist
    topics = (
        db.query(Topic)
        .filter(Topic.topic.in_([topic.topic for topic in survey_request.topics]))
        .all()
    )
    if not topics:
        raise HTTPException(status_code=404, detail="Topics not found")
    # Initialize channel and delivery_service to None
    channel = None
    delivery_service = None

    # If channel is provided, check if channel exists
    if survey_request.channel:
        channel = (
            db.query(Channel).filter(Channel.name == survey_request.channel).first()
        )
        if not channel:
            raise HTTPException(status_code=404, detail="Channel not found")
    # If delivery service is provided, check if delivery service exists
    if survey_request.delivery_service:
        delivery_service = (
            db.query(DeliveryService)
            .filter(DeliveryService.name == survey_request.delivery_service)
            .first()
        )
        if not delivery_service:
            raise HTTPException(status_code=404, detail="Delivery service not found")
    try:
        # Create survey without departments, topics, and keywords relationships
        survey = Survey(
            store_key=survey_request.store_key,
            comment=survey_request.comment,
            sentiment=survey_request.sentiment,
            reported_at=as_utc(survey_request.reported_at),
            cls=survey_request.cls,
            channel_id=channel.id if channel else None,
            delivery_service_id=delivery_service.id if delivery_service else None,
        )
        db.add(survey)
        db.flush()
        # Add departments, topics, and keywords relationships
        for department in departments:
            survey_department = SurveyDepartments(
                survey_id=survey.id,
                department_id=department.id,
                sentiment=survey_request.sentiment,
            )
            db.add(survey_department)
        for request_topic in survey_request.topics:
            topic = db.query(Topic).filter(Topic.topic == request_topic.topic).first()
            survey_topic = SurveyTopics(
                survey_id=survey.id,
                topic_id=topic.id,
                sentiment=request_topic.sentiment,
            )
            db.add(survey_topic)
        for request_keyword in survey_request.keywords:
            keyword = (
                db.query(Keyword)
                .filter(Keyword.keyword == request_keyword.keyword)
                .first()
            )
            # Create keyword if it doesn't exist
            if not keyword:
                keyword = Keyword(keyword=request_keyword.keyword)
                db.add(keyword)
                db.flush()
            survey_keyword = SurveyKeywords(
                survey_id=survey.id,
                keyword_id=keyword.id,
                sentiment=request_keyword.sentiment,
            )
            db.add(survey_keyword)
        db.commit()
        db.refresh(survey)
        return survey.to_dict()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/download")
async def download_surveys(
    filter_params: FilterRequest = Depends(get_filter_params),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    filter_dict = filter_params.model_dump()
    
    # Get distinct survey IDs that match the filters
    id_query = build_survey_query(db.query(Survey.id).distinct(), filter_dict)
    survey_ids = [row[0] for row in id_query.all()]
    
    def format_excel_value(value):
        """Format a value for Excel export, handling None, lists, dates, enums, and strings."""
        if value is None:
            return ""
        elif isinstance(value, list):
            return "; ".join(format_excel_value(item) for item in value)
        elif isinstance(value, datetime):
            return utc_isoformat(value) or ""
        elif hasattr(value, 'value'):  # Handle enum types
            # Extract the enum value and capitalize: "POSITIVE" -> "Positive"
            enum_value = value.value
            return enum_value.title() if isinstance(enum_value, str) else str(enum_value)
        elif isinstance(value, str):
            # Handle string representations of enums like "Sentiment.NEUTRAL" or "POSITIVE"
            if '.' in value and any(enum_name in value for enum_name in ['Sentiment', 'TopicSentiment']):
                # Extract just the value part: "Sentiment.NEUTRAL" -> "NEUTRAL" -> "Neutral"
                enum_value = value.split('.')[-1]
                return enum_value.title()
            return value
        else:
            return str(value)

    def generate_excel_file():
        # Create a workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Surveys"

        # Define headers
        headers = [
            "id",
            "survey_id",
            "respondent_id",
            "store_key",
            "store_name_english",
            "store_name_local",
            "bu_key",
            "area_manager",
            "store_format",
            "store_type",
            "operations_controller",
            "regional_manager",
            "px",
            "csr",
            "dr",
            "mag_type",
            "cf_grouping",
            "store_brand",
            "competitor",
            "region",
            "area",
            "province",
            "territory",
            "toh",
            "district",
            "city",
            "operations_manager",
            "district_manager",
            "sic",
            "soc",
            "tech_life_type",
            "operation_manager_tl",
            "region_manager_tl",
            "relocation",
            "latitude",
            "longitude",
            "store_open_date",
            "store_close_date",
            "is_closed",
            #"department_id",
            "department_name",
            #"channel_id",
            "channel_name",
            "departments",
            "topics",
            "keywords",
            "comment",
            "channel",
            "delivery_service",
            "sentiment", # topic_sentiment
            "sentiment_score", # topic_sentiment_score
            "cls",
            "reported_at",
            "created_at",
            "updated_at",
        ]

        export_headers = [
            header for header in headers if is_survey_export_column_enabled(header)
        ]
        if not export_headers:
            raise HTTPException(
                status_code=500,
                detail="No survey export columns enabled. Set SURVEY_EXPORT_COLUMN_<COLUMN_NAME>=true for required columns.",
            )

        # Write headers to first row
        for col_idx, header in enumerate(export_headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Stream surveys in batches using offset
        batch_size = 100
        offset = 0
        row_num = 2  # Start from row 2 (row 1 is headers)

        while True:
            # Get surveys for the current batch of IDs
            batch_ids = survey_ids[offset:offset + batch_size]
            if not batch_ids:
                break
                
            surveys = (
                db.query(Survey)
                .filter(Survey.id.in_(batch_ids))
                .options(
                    joinedload(Survey.store),
                    joinedload(Survey.survey_topics),
                    joinedload(Survey.survey_keywords),
                    joinedload(Survey.survey_departments),
                    joinedload(Survey.channel),
                    joinedload(Survey.delivery_service)
                )
                .order_by(Survey.reported_at.asc())
                .all()
            )
            for survey in surveys:
                csv_value = survey.to_csv(filter_params.timezone)
                # Write row values
                for col_idx, header in enumerate(export_headers, start=1):
                    value = format_excel_value(csv_value.get(header))
                    ws.cell(row=row_num, column=col_idx, value=value)
                row_num += 1
            offset += batch_size

        # Save workbook to BytesIO buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    excel_buffer = await asyncio.to_thread(generate_excel_file)
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=surveys.xlsx"},
    )


@router.get("/{survey_id}")
async def get_survey(
    survey_id: int,
    timezone: Optional[str] = Query(
        default=None,
        description="Optional IANA timezone for timestamp display; UTC is used when omitted",
    ),
    db: Session = Depends(get_db),
) -> SurveyResponse:
    try:
        resolve_timezone(timezone)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    filter_dict = {"ids": [survey_id]}
    query = build_survey_query(db.query(Survey), filter_dict)
    query = query.options(
        joinedload(Survey.store),
        joinedload(Survey.survey_topics),
        joinedload(Survey.survey_keywords),
        joinedload(Survey.survey_departments),
        joinedload(Survey.channel),
        joinedload(Survey.delivery_service)
    )
    survey = query.first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found")
    return SurveyResponse.model_validate(survey.to_dict(timezone))


class UpdateSurveyTopicRequest(BaseModel):
    topic: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyKeywordRequest(BaseModel):
    keyword: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyDepartmentRequest(BaseModel):
    name: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyRequest(BaseModel):
    store_key: Optional[int] = None
    channel: Optional[str] = None
    delivery_service: Optional[str] = None
    departments: Optional[List[UpdateSurveyDepartmentRequest]] = None
    comment: Optional[str] = None
    sentiment: Optional[Literal["positive", "negative", "neutral"]] = None
    topics: Optional[List[UpdateSurveyTopicRequest]] = None
    keywords: Optional[List[UpdateSurveyKeywordRequest]] = None
    reported_at: Optional[datetime] = None
    cls: Optional[float] = None


@router.put("/{survey_id}")
async def update_survey(
    survey_id: int,
    survey_request: UpdateSurveyRequest,
    db: Session = Depends(get_db),
    _: object = Depends(require_admin),
) -> SurveyResponse:
    try:
        filter_dict = {"ids": [survey_id]}
        query = build_survey_query(db.query(Survey), filter_dict)
        query = query.options(
            joinedload(Survey.store),
            joinedload(Survey.survey_topics),
            joinedload(Survey.survey_keywords),
            joinedload(Survey.survey_departments),
            joinedload(Survey.channel),
            joinedload(Survey.delivery_service)
        )
        survey = query.first()
        if not survey:
            raise HTTPException(status_code=404, detail="Survey not found")
        if survey_request.store_key:
            survey.store_key = survey_request.store_key
        if survey_request.channel:
            survey.channel_id = survey_request.channel
        if survey_request.delivery_service:
            survey.delivery_service_id = survey_request.delivery_service
        if survey_request.departments:
            for department in survey_request.departments:
                survey_department = (
                    db.query(SurveyDepartments)
                    .filter(
                        SurveyDepartments.survey_id == survey.id,
                        SurveyDepartments.department_id == department.name,
                    )
                    .first()
                )
                if not survey_department:
                    survey_department = SurveyDepartments(
                        survey_id=survey.id,
                        department_id=department.name,
                        sentiment=department.sentiment,
                    )
                    db.add(survey_department)
                else:
                    survey_department.sentiment = department.sentiment
        if survey_request.comment:
            survey.comment = survey_request.comment
        if survey_request.sentiment:
            survey.sentiment = survey_request.sentiment
        if survey_request.topics:
            for topic in survey_request.topics:
                survey_topic = (
                    db.query(SurveyTopics)
                    .filter(
                        SurveyTopics.survey_id == survey.id,
                        SurveyTopics.topic_id == topic.topic,
                    )
                    .first()
                )
                if not survey_topic:
                    survey_topic = SurveyTopics(
                        survey_id=survey.id,
                        topic_id=topic.topic,
                        sentiment=topic.sentiment,
                    )
                    db.add(survey_topic)
                else:
                    survey_topic.sentiment = topic.sentiment
        if survey_request.keywords:
            for keyword in survey_request.keywords:
                survey_keyword = (
                    db.query(SurveyKeywords)
                    .filter(
                        SurveyKeywords.survey_id == survey.id,
                        SurveyKeywords.keyword_id == keyword.keyword,
                    )
                    .first()
                )
                if not survey_keyword:
                    survey_keyword = SurveyKeywords(
                        survey_id=survey.id,
                        keyword_id=keyword.keyword,
                        sentiment=keyword.sentiment,
                    )
                    db.add(survey_keyword)
                else:
                    survey_keyword.sentiment = keyword.sentiment
        if survey_request.reported_at:
            survey.reported_at = as_utc(survey_request.reported_at)
        if survey_request.cls is not None:
            survey.cls = survey_request.cls
        db.commit()
        db.refresh(survey)
        return SurveyResponse.model_validate(survey.to_dict())
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/{survey_id}")
async def delete_survey(
    survey_id: int,
    db: Session = Depends(get_db),
    _: object = Depends(require_admin),
) -> SurveyResponse:
    try:
        filter_dict = {"ids": [survey_id]}
        query = build_survey_query(db.query(Survey), filter_dict)
        query = query.options(
            joinedload(Survey.store),
            joinedload(Survey.survey_topics),
            joinedload(Survey.survey_keywords),
            joinedload(Survey.survey_departments),
            joinedload(Survey.channel),
            joinedload(Survey.delivery_service)
        )
        survey = query.first()
        if not survey:
            raise HTTPException(status_code=404, detail="Survey not found")
        survey.is_deleted = True
        db.commit()
        db.refresh(survey)
        return SurveyResponse.model_validate(survey.to_dict())
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
