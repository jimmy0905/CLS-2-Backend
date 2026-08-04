from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy.orm import Session, joinedload
from utils.database import get_db
from models.User import User
from utils.conditionFilter import build_survey_query, build_optimized_query
from utils.security import get_current_user
from models.Survey import Survey
from typing import List, Optional
from utils.llm.generate_strategy import (
    generate_store_strategy,
    generate_region_strategy,
    generate_channel_strategy,
    generate_delivery_service_strategy,
)
from pydantic import BaseModel, Field
from models.Store import Store
from models.Channel import Channel
from models.DeliveryService import DeliveryService
from sqlalchemy import or_, func, case
from utils.conditionFilter import FilterRequest, get_filter_params
from models.enum.Sentiment import TopicSentiment, Sentiment
from datetime import datetime, date
from io import BytesIO
import requests
from openpyxl import Workbook
import os
import asyncio
from fastapi import BackgroundTasks
from utils.utc import utc_isoformat


router = APIRouter(
    prefix="/strategy",
    tags=["strategy"],
    dependencies=[Depends(get_current_user)],
)

class StoreResponse(BaseModel):
    store_key: int
    store_name_english: Optional[str] = None
    store_name_local: Optional[str] = None
    bu_key: Optional[str] = None
    area_manager: Optional[str] = None
    store_format: Optional[str] = None
    store_type: Optional[str] = None
    operations_controller: Optional[str] = None
    regional_manager: Optional[str] = None
    px: Optional[str] = None
    csr: Optional[str] = None
    dr: Optional[str] = None
    mag_type: Optional[str] = None
    cf_grouping: Optional[str] = None
    store_brand: Optional[str] = None
    competitor: Optional[str] = None
    region: Optional[str] = None
    area: Optional[str] = None
    territory: Optional[str] = None
    toh: Optional[str] = None
    district: Optional[str] = None
    city: Optional[str] = None
    operations_manager: Optional[str] = None
    district_manager: Optional[str] = None
    sic: Optional[str] = None
    soc: Optional[str] = None
    tech_life_type: Optional[str] = None
    operation_manager_tl: Optional[str] = None
    region_manager_tl: Optional[str] = None
    relocation: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    store_open_date: Optional[date] = None
    store_close_date: Optional[date] = None
    is_closed: bool


class TopKPerformanceStoresResponse(BaseModel):
    store: StoreResponse
    score: float
    positive_count: int
    negative_count: int
    neutral_count: int
    mixed_count: int


@router.get("/get_top_k_performance_stores")
async def get_top_k_performance_stores(
    db: Session = Depends(get_db),
    filter_params: FilterRequest = Depends(get_filter_params),
    k: int = Query(
        default=10,
        description="The number of top performing stores to return",
    ),
) -> List[TopKPerformanceStoresResponse]:
    filter_dict = filter_params.model_dump()
    # First get the sentiment counts per store
    sentiment_query, sentiment_joins, _ = build_optimized_query(db, filter_dict)
    if "store" not in sentiment_joins:
        sentiment_query = sentiment_query.join(Store, Survey.store_key == Store.store_key)

    sentiment_results = (
        sentiment_query.with_entities(
            Store.store_key.label("store_key"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.NEUTRAL, Survey.id))
            ).label("neutral_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.POSITIVE, Survey.id))
            ).label("positive_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.NEGATIVE, Survey.id))
            ).label("negative_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.MIXED, Survey.id))
            ).label("mixed_count"),
            func.avg(Survey.topic_sentiment_score).label("average_sentiment_score"),
        )
        .group_by(Store.store_key)
        .all()
    )

    # Get store details for the stores we found
    store_keys = [r.store_key for r in sentiment_results]
    stores = (
        db.query(Store)
        .filter(Store.store_key.in_(store_keys))
        .all()
    )

    # Create a lookup dictionary for store details
    store_details_map = {s.store_key: s for s in stores}

    store_with_sentiment = []
    for store in sentiment_results:
        store_obj = store_details_map[store.store_key]
        average_sentiment_score = store.average_sentiment_score if store.average_sentiment_score is not None else 0
        store_with_sentiment.append(
            TopKPerformanceStoresResponse(
                store={
                    "store_key": store.store_key,
                    "store_name_english": store_obj.store_name_english,
                    "store_name_local": store_obj.store_name_local,
                    "bu_key": store_obj.bu_key,
                    "area_manager": store_obj.area_manager,
                    "store_format": store_obj.store_format,
                    "store_type": store_obj.store_type,
                    "operations_controller": store_obj.operations_controller,
                    "regional_manager": store_obj.regional_manager,
                    "px": store_obj.px,
                    "csr": store_obj.csr,
                    "dr": store_obj.dr,
                    "mag_type": store_obj.mag_type,
                    "cf_grouping": store_obj.cf_grouping,
                    "store_brand": store_obj.store_brand,
                    "competitor": store_obj.competitor,
                    "region": store_obj.region,
                    "area": store_obj.area,
                    "territory": store_obj.territory,
                    "toh": store_obj.toh,
                    "district": store_obj.district,
                    "city": store_obj.city,
                    "operations_manager": store_obj.operations_manager,
                    "district_manager": store_obj.district_manager,
                    "sic": store_obj.sic,
                    "soc": store_obj.soc,
                    "tech_life_type": store_obj.tech_life_type,
                    "operation_manager_tl": store_obj.operation_manager_tl,
                    "region_manager_tl": store_obj.region_manager_tl,
                    "relocation": store_obj.relocation,
                    "latitude": store_obj.latitude,
                    "longitude": store_obj.longitude,
                    "store_open_date": store_obj.store_open_date,
                    "store_close_date": store_obj.store_close_date,
                    "is_closed": store_obj.is_closed,
                },
                score=average_sentiment_score,
                positive_count=store.positive_count,
                negative_count=store.negative_count,
                neutral_count=store.neutral_count,
                mixed_count=store.mixed_count,
            )
        )
    store_with_sentiment.sort(key=lambda x: x.score, reverse=True)
    store_with_sentiment = store_with_sentiment[:k]
    return store_with_sentiment


class StrategyByStoreIdsRequest(BaseModel):
    store_keys: List[str] = Field(
        default_factory=list,
        description="The store ids to get the strategy",
    )


@router.post("/get_strategy_for_store_by_ids")
async def get_strategy_for_store_by_ids(
    request: StrategyByStoreIdsRequest,
    db: Session = Depends(get_db),
) -> str:
    filter_dict = {
        "store_keys": request.store_keys,
        "topic_sentiments": [TopicSentiment.POSITIVE],
    }
    filtered_query, _, _ = build_optimized_query(db, filter_dict)
    surveys = (
        filtered_query.order_by(func.length(Survey.comment).desc()).limit(30).all()
    )
    strategy, _ = await generate_store_strategy(surveys)
    return strategy


class TopKPerformanceColumnResponse(BaseModel):
    column_value: str
    score: float
    positive_count: int
    negative_count: int
    neutral_count: int
    mixed_count: int


@router.get("/get_top_k_performance_columns")
async def get_top_k_performance_columns(
    column: str = Query(..., description="The column name to get the sentiment distribution for store column, columns are bu_key, area_manager, store_format, store_type, operations_controller, regional_manager, px, csr, dr, mag_type, cf_grouping, store_brand, competitor, region, area, province, territory, toh, district, city, operations_manager, district_manager, sic, soc, tech_life_type, operation_manager_tl, region_manager_tl, relocation, latitude, longitude, store_open_date, store_close_date, is_closed, store_key, store_english_name, store_local_name"),
    db: Session = Depends(get_db),
    filter_params: FilterRequest = Depends(get_filter_params),
    k: int = Query(
        default=10,
        description="The number of top performing columns to return",
    ),
) -> List[TopKPerformanceColumnResponse]:

    column_name_mapper = {
        "bu_key": Store.bu_key,
        "area_manager": Store.area_manager,
        "store_format": Store.store_format,
        "store_type": Store.store_type,
        "operations_controller": Store.operations_controller,
        "regional_manager": Store.regional_manager,
        "px": Store.px,
        "csr": Store.csr,
        "dr": Store.dr,
        "mag_type": Store.mag_type,
        "cf_grouping": Store.cf_grouping,
        "store_brand": Store.store_brand,
        "competitor": Store.competitor,
        "region": Store.region,
        "area": Store.area,
        "province": Store.province,
        "territory": Store.territory,
        "toh": Store.toh,
        "district": Store.district,
        "city": Store.city,
        "operations_manager": Store.operations_manager,
        "district_manager": Store.district_manager,
        "sic": Store.sic,
        "soc": Store.soc,
        "tech_life_type": Store.tech_life_type,
        "operation_manager_tl": Store.operation_manager_tl,
        "region_manager_tl": Store.region_manager_tl,
        "relocation": Store.relocation,
        "latitude": Store.latitude,
        "longitude": Store.longitude,
        "store_open_date": Store.store_open_date,
        "store_close_date": Store.store_close_date,
        "is_closed": Store.is_closed,
        "store_key": Store.store_key,
        "store_english_name": Store.store_name_english,
        "store_local_name": Store.store_name_local,
    }
    
    column_col = column_name_mapper.get(column)
    if column_col is None:
        raise HTTPException(status_code=404, detail="Column name not found")

    filter_dict = filter_params.model_dump()


    sentiment_query, sentiment_joins, _ = build_optimized_query(
        db, filter_dict
    )
    if "store" not in sentiment_joins:
        sentiment_query = sentiment_query.join(Store, Survey.store_key == Store.store_key)

    sentiment_results = (
        sentiment_query.with_entities(
            column_col.label("column_name"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.NEUTRAL, Survey.id))
            ).label("neutral_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.POSITIVE, Survey.id))
            ).label("positive_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.NEGATIVE, Survey.id))
            ).label("negative_count"),
            func.count(
                case((Survey.topic_sentiment == TopicSentiment.MIXED, Survey.id))
            ).label("mixed_count"),
            func.avg(Survey.topic_sentiment_score).label("average_sentiment_score"),
        )
        .group_by(column_col)
        .all()
    )
    hierarchy_with_sentiment = []
    for hierarchy in sentiment_results:
        average_sentiment_score = hierarchy.average_sentiment_score if hierarchy.average_sentiment_score is not None else 0
        hierarchy_with_sentiment.append(
            TopKPerformanceColumnResponse(
                column_value=str(hierarchy.column_name),
                score=average_sentiment_score,
                positive_count=hierarchy.positive_count,
                negative_count=hierarchy.negative_count,
                neutral_count=hierarchy.neutral_count,
                mixed_count=hierarchy.mixed_count,
            )
        )
    hierarchy_with_sentiment.sort(key=lambda x: x.score, reverse=True)
    hierarchy_with_sentiment = hierarchy_with_sentiment[:k]
    return hierarchy_with_sentiment

class StrategyByColumnValueRequest(BaseModel):
    target_column: str
    column_values: List[str]

@router.post("/get_strategy_for_column_by_values")
async def get_strategy_for_column_by_value(
    request: StrategyByColumnValueRequest,
    db: Session = Depends(get_db),
) -> str:
    if not request.column_values:
        raise HTTPException(
            status_code=400,
            detail="column_values must contain at least one value",
        )

    column_name_mapper = {
        "bu_key": Store.bu_key,
        "area_manager": Store.area_manager,
        "store_format": Store.store_format,
        "store_type": Store.store_type,
        "operations_controller": Store.operations_controller,
        "regional_manager": Store.regional_manager,
        "px": Store.px,
        "csr": Store.csr,
        "dr": Store.dr,
        "mag_type": Store.mag_type,
        "cf_grouping": Store.cf_grouping,
        "store_brand": Store.store_brand,
        "competitor": Store.competitor,
        "region": Store.region,
        "area": Store.area,
        "province": Store.province,
        "territory": Store.territory,
        "toh": Store.toh,
        "district": Store.district,
        "city": Store.city,
        "operations_manager": Store.operations_manager,
        "district_manager": Store.district_manager,
        "sic": Store.sic,
        "soc": Store.soc,
        "tech_life_type": Store.tech_life_type,
        "operation_manager_tl": Store.operation_manager_tl,
        "region_manager_tl": Store.region_manager_tl,
        "relocation": Store.relocation,
        "latitude": Store.latitude,
        "longitude": Store.longitude,
        "store_open_date": Store.store_open_date,
        "store_close_date": Store.store_close_date,
        "is_closed": Store.is_closed,
        "store_key": Store.store_key,
        "store_english_name": Store.store_name_english,
        "store_local_name": Store.store_name_local,
    }
    column_to_filter_key_mapper = {
        "bu_key": "bu_keys",
        "area_manager": "area_managers",
        "store_format": "store_formats",
        "store_type": "store_types",
        "operations_controller": "operations_controllers",
        "regional_manager": "regional_managers",
        "px": "px",
        "csr": "csr",
        "dr": "dr",
        "mag_type": "mag_types",
        "cf_grouping": "cf_groupings",
        "store_brand": "store_brands",
        "competitor": "competitors",
        "region": "regions",
        "area": "areas",
        "province": "provinces",
        "territory": "territories",
        "toh": "tohs",
        "district": "districts",
        "city": "cities",
        "operations_manager": "operations_managers",
        "district_manager": "district_managers",
        "sic": "sic",
        "soc": "soc",
        "tech_life_type": "tech_life_types",
        "operation_manager_tl": "operation_manager_tls",
        "region_manager_tl": "region_manager_tls",
        "relocation": "relocations",
        "latitude": "latitudes",
        "longitude": "longitudes",
        "store_open_date": "store_open_dates",
        "store_close_date": "store_close_dates",
        "is_closed": "is_closed",
        "store_key": "store_keys",
        "store_english_name": "store_english_names",
        "store_local_name": "store_local_names",
    }

    column_col = column_name_mapper.get(request.target_column)
    if column_col is None:
        raise HTTPException(status_code=404, detail="Column name not found")
    filter_key = column_to_filter_key_mapper.get(request.target_column)
    if filter_key is None:
        raise HTTPException(status_code=404, detail="Filter key not found for column")

    filter_value: list[str] | list[int] | list[float] | list[date] | bool = request.column_values
    if filter_key == "is_closed":
        if len(request.column_values) != 1:
            raise HTTPException(
                status_code=400,
                detail="is_closed requires exactly one boolean value",
            )

        normalized_value = request.column_values[0].strip().lower()
        if normalized_value not in {"true", "false"}:
            raise HTTPException(
                status_code=400,
                detail="is_closed must be 'true' or 'false'",
            )
        filter_value = normalized_value == "true"
    elif filter_key == "store_keys":
        try:
            filter_value = [int(value) for value in request.column_values]
        except ValueError as exc:
            raise HTTPException(
                status_code=400,
                detail="store_key values must be integers",
            ) from exc
    elif filter_key in {"latitudes", "longitudes"}:
        try:
            filter_value = [float(value) for value in request.column_values]
        except ValueError as exc:
            raise HTTPException(
                status_code=400,
                detail=f"{request.target_column} values must be numbers",
            ) from exc
    elif filter_key in {"store_open_dates", "store_close_dates"}:
        try:
            filter_value = [date.fromisoformat(value) for value in request.column_values]
        except ValueError as exc:
            raise HTTPException(
                status_code=400,
                detail=f"{request.target_column} values must use YYYY-MM-DD format",
            ) from exc

    filter_dict = {
        filter_key: filter_value,
        "topic_sentiments": [TopicSentiment.POSITIVE],
    }
    filtered_query, _, _ = build_optimized_query(db, filter_dict)
    surveys = (
        filtered_query.order_by(func.length(Survey.comment).desc()).limit(30).all()
    )
    strategy, _ = await generate_region_strategy(surveys)
    return strategy

class ChannelResponse(BaseModel):
    id: int
    name: str

class TopKPerformanceChannelsResponse(BaseModel):
    channel: ChannelResponse
    score: float
    positive_count: int
    negative_count: int
    neutral_count: int
    mixed_count: int


@router.get("/get_top_k_performance_channels")
async def get_top_k_performance_channels(
    db: Session = Depends(get_db),
    filter_params: FilterRequest = Depends(get_filter_params),
    k: int = Query(
        default=10,
        description="The number of top performing channels to return",
    ),
) -> List[TopKPerformanceChannelsResponse]:
    filter_dict = filter_params.model_dump()
    sentiment_query, sentiment_joins, _ = build_optimized_query(db, filter_dict)
    if "store" not in sentiment_joins:
        sentiment_query = sentiment_query.join(Store, Survey.store_key == Store.store_key)
    
    # Add Channel join if not already present
    if "channel" not in sentiment_joins:
        sentiment_query = sentiment_query.join(Channel, Survey.channel_id == Channel.id)

    sentiment_results = (
        sentiment_query.with_entities(
            Channel.id.label("channel_id"),
            Channel.name.label("channel_name"),
            func.count(case((Survey.topic_sentiment == TopicSentiment.NEUTRAL, Survey.id))).label(
                "neutral_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.POSITIVE, Survey.id))).label(
                "positive_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.NEGATIVE, Survey.id))).label(
                "negative_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.MIXED, Survey.id))).label(
                "mixed_count"
            ),
            func.avg(Survey.topic_sentiment_score).label("average_sentiment_score"),
        )
        .group_by(Channel.id, Channel.name)
        .all()
    )
    channel_with_sentiment: List[TopKPerformanceChannelsResponse] = []
    for channel in sentiment_results:
        average_sentiment_score = channel.average_sentiment_score if channel.average_sentiment_score is not None else 0
        channel_with_sentiment.append(
            TopKPerformanceChannelsResponse(
                channel={
                    "id": channel.channel_id,
                    "name": channel.channel_name,
                },
                score=average_sentiment_score,
                positive_count=channel.positive_count,
                negative_count=channel.negative_count,
                neutral_count=channel.neutral_count,
                mixed_count=channel.mixed_count,
            )
        )
    channel_with_sentiment.sort(key=lambda x: x.score, reverse=True)
    channel_with_sentiment = channel_with_sentiment[:k]
    return channel_with_sentiment

class TopKPerformanceByChannelIdsRequest(BaseModel):
    channel_ids: List[int] = Field(
        default_factory=list,
        description="The channel ids to get the top performing channels",
    )

@router.post("/get_strategy_for_channel_by_ids")
async def get_top_k_performance_channels_by_ids(
    request: TopKPerformanceByChannelIdsRequest,
    db: Session = Depends(get_db),
) -> str:
    filter_dict = {
        "channel_ids": request.channel_ids,
        "topic_sentiments": [TopicSentiment.POSITIVE],
    }
    filtered_query, _, _ = build_optimized_query(db, filter_dict)
    surveys = (
        filtered_query.order_by(func.length(Survey.comment).desc()).limit(30).all()
    )
    strategy, _ = await generate_channel_strategy(surveys)
    return strategy

class DeliveryServiceResponse(BaseModel):
    id: int
    name: str

class TopKPerformanceDeliveryServicesResponse(BaseModel):
    delivery_service: DeliveryServiceResponse
    score: float
    positive_count: int
    negative_count: int
    neutral_count: int
    mixed_count: int

@router.get("/get_top_k_performance_delivery_services")
async def get_top_k_performance_delivery_services(
    db: Session = Depends(get_db),
    filter_params: FilterRequest = Depends(get_filter_params),
    k: int = Query(
        default=10,
        description="The number of top performing delivery services to return",
    ),
) -> List[TopKPerformanceDeliveryServicesResponse]:
    filter_dict = filter_params.model_dump()
    sentiment_query, sentiment_joins, _ = build_optimized_query(db, filter_dict)
    if "store" not in sentiment_joins:
        sentiment_query = sentiment_query.join(Store, Survey.store_key == Store.store_key)
    
    # Add DeliveryService join if not already present
    if "delivery_service" not in sentiment_joins:
        sentiment_query = sentiment_query.join(DeliveryService, Survey.delivery_service_id == DeliveryService.id)

    sentiment_results = (
        sentiment_query.with_entities(
            DeliveryService.id.label("delivery_service_id"),
            DeliveryService.name.label("delivery_service_name"),
            func.count(case((Survey.topic_sentiment == TopicSentiment.NEUTRAL, Survey.id))).label(
                "neutral_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.POSITIVE, Survey.id))).label(
                "positive_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.NEGATIVE, Survey.id))).label(
                "negative_count"
            ),
            func.count(case((Survey.topic_sentiment == TopicSentiment.MIXED, Survey.id))).label(
                "mixed_count"
            ),
            func.avg(Survey.topic_sentiment_score).label("average_sentiment_score"),
        )
        .group_by(DeliveryService.id, DeliveryService.name)
        .all()
    )
    delivery_service_with_sentiment: List[TopKPerformanceDeliveryServicesResponse] = []
    for delivery_service in sentiment_results:
        average_sentiment_score = delivery_service.average_sentiment_score if delivery_service.average_sentiment_score is not None else 0
        delivery_service_with_sentiment.append(
            TopKPerformanceDeliveryServicesResponse(
                delivery_service={
                    "id": delivery_service.delivery_service_id,
                    "name": delivery_service.delivery_service_name,
                },
                score=average_sentiment_score,
                positive_count=delivery_service.positive_count,
                negative_count=delivery_service.negative_count,
                neutral_count=delivery_service.neutral_count,
                mixed_count=delivery_service.mixed_count,
            )
        )
    delivery_service_with_sentiment.sort(key=lambda x: x.score, reverse=True)
    delivery_service_with_sentiment = delivery_service_with_sentiment[:k]
    return delivery_service_with_sentiment 

class TopKPerformanceByDeliveryServiceIdsRequest(BaseModel):
    delivery_service_ids: List[int] = Field(
        default_factory=list,
        description="The delivery service ids to get the top performing delivery services",
    )

@router.post("/get_strategy_for_delivery_service_by_ids")
async def get_top_k_performance_delivery_services_by_ids(
    request: TopKPerformanceByDeliveryServiceIdsRequest,
    db: Session = Depends(get_db),
) -> str:
    filter_dict = {
        "delivery_service_ids": request.delivery_service_ids,
        "topic_sentiments": [TopicSentiment.POSITIVE],
    }
    filtered_query, _, _ = build_optimized_query(db, filter_dict)
    surveys = (
        filtered_query.order_by(func.length(Survey.comment).desc()).limit(30).all()
    )
    strategy, _ = await generate_delivery_service_strategy(surveys)
    return strategy

@router.get("/get_strategy_v2")
async def get_strategy_v2(
    db: Session = Depends(get_db),
    filter_params: FilterRequest = Depends(get_filter_params),
) -> dict:
    filter_dict = filter_params.model_dump()
    
    # Get distinct survey IDs that match the filters
    id_query = build_survey_query(db.query(Survey.id).distinct(), filter_dict)
    survey_ids = [row[0] for row in id_query.all()]
    
    def format_excel_value(value):
        from datetime import date
        if value is None:
            return ""
        elif isinstance(value, list):
            return "; ".join(str(item) for item in value)
        elif isinstance(value, datetime):
            return utc_isoformat(value) or ""
        elif isinstance(value, date):
            return value.isoformat()
        else:
            return str(value)

    def generate_excel_and_call_api():
        # Create a workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Surveys"

        # Define headers
        headers = [
            "id",
            "survey_id",
            "respondent_id",
            "store_key",
            "store_id",
            "store_name",
            "store_name_english",
            "store_name_local",
            "bu_key",
            "area_manager",
            "store_format",
            "store_type",
            "operations_controller",
            "regional_manager",
            "px",
            "csr",
            "dr",
            "mag_type",
            "cf_grouping",
            "store_brand",
            "competitor",
            "region",
            "area",
            "territory",
            "toh",
            "district",
            "city",
            "operations_manager",
            "district_manager",
            "sic",
            "soc",
            "tech_life_type",
            "operation_manager_tl",
            "region_manager_tl",
            "relocation",
            "latitude",
            "longitude",
            "store_open_date",
            "store_close_date",
            "is_closed",
            "department_id",
            "department_name",
            "channel_id",
            "channel_name",
            "departments",
            "topics",
            "keywords",
            "comment",
            "channel",
            "delivery_service",
            "sentiment", # topic_sentiment
            "sentiment_score", # topic_sentiment_score
            "reported_at",
            "created_at",
            "updated_at",
        ]

        # Write headers to first row
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Stream surveys in batches using offset
        batch_size = 100
        offset = 0
        row_num = 2  # Start from row 2 (row 1 is headers)

        while True:
            # Get surveys for the current batch of IDs
            batch_ids = survey_ids[offset:offset + batch_size]
            if not batch_ids:
                break
                
            surveys = (
                db.query(Survey)
                .filter(Survey.id.in_(batch_ids))
                .options(
                    joinedload(Survey.store),
                    joinedload(Survey.survey_topics),
                    joinedload(Survey.survey_keywords),
                    joinedload(Survey.survey_departments),
                    joinedload(Survey.channel),
                    joinedload(Survey.delivery_service)
                )
                .order_by(Survey.reported_at.asc())
                .all()
            )
            for survey in surveys:
                csv_value = survey.to_csv()
                # Write row values
                for col_idx, header in enumerate(headers, start=1):
                    value = format_excel_value(csv_value[header])
                    ws.cell(row=row_num, column=col_idx, value=value)
                row_num += 1
            offset += batch_size

        # Save workbook to BytesIO buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # Make API request
        url = os.getenv("ANALYZE_FEEDBACK_API_URL")
        if not url:
            raise Exception("ANALYZE_FEEDBACK_API_URL is not set")
        
        url = url.rstrip('/') + "/analyze-feedback"
        # payload = {'analysis_mode': 'STAT',
        # 'sampling_method': 'DIRECT',
        # 'top_n_stores': '10',
        # 'bottom_n_stores': '10',
        # 'top_n_topics': '10',
        # 'quote_sample_size': '5'}
        payload = {
            'include_channel': os.getenv("ANALYZE_FEEDBACK_IS_INCLUDE_CHANNEL") == "true",
        }
        files=[
        ('file',('surveys.xlsx',buffer,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        ]
        
        request_headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
        }

        session = requests.Session()
        session.trust_env = False
        response = session.post(url, data=payload, files=files, headers=request_headers, timeout=300)
        
        if response.status_code != 200:
            raise Exception(f"Analysis API error: {response.text[:500]}")
        
        return response.json()
    
    result = await asyncio.to_thread(generate_excel_and_call_api)
    return result