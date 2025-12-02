from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import and_, func, or_, select
from models.Survey import Survey
from models.Topic import Topic
from models.Keyword import Keyword
from models.SurveyTopics import SurveyTopics
from models.SurveyKeywords import SurveyKeywords
from models.SurveyDepartments import SurveyDepartments
from models.Store import Store
from models.Hierarchy import Hierarchy
from models.Department import Department
from models.Channel import Channel
from models.DeliveryService import DeliveryService
from utils.database import get_db
from pydantic import BaseModel, Field
from typing import Literal, List, Optional
from datetime import datetime
from utils.conditionFilter import (
    build_survey_query,
    FilterRequest,
    get_filter_params,
)
from utils.llm.extract_total import (
    extract_total,
    extract_total_retry,
)
from utils.security import get_current_user
from fastapi_pagination import Page, paginate
from fastapi.responses import StreamingResponse
import logging
from openpyxl import Workbook
from io import BytesIO

logger = logging.getLogger(__name__)


router = APIRouter(
    prefix="/surveys",
    tags=["surveys"],
    dependencies=[Depends(get_db), Depends(get_current_user)],
)


class HierarchyResponse(BaseModel):
    id: int
    name: str
    level: int


class StoreResponse(BaseModel):
    id: int
    name: str
    hierarchy_level_1: Optional[HierarchyResponse] = None
    hierarchy_level_2: Optional[HierarchyResponse] = None
    hierarchy_level_3: Optional[HierarchyResponse] = None
    hierarchy_level_4: Optional[HierarchyResponse] = None
    hierarchy_level_5: Optional[HierarchyResponse] = None


class ChannelResponse(BaseModel):
    id: int
    name: str


class DeliveryServiceResponse(BaseModel):
    id: int
    name: str


class DepartmentWithSentimentResponse(BaseModel):
    department_id: int
    name: str
    sentiment: str


class KeywordWithSentimentResponse(BaseModel):
    keyword_id: int
    keyword: str
    sentiment: str


class TopicWithSentimentResponse(BaseModel):
    topic_id: int
    topic: str
    sentiment: str


class SurveyResponse(BaseModel):
    id: int
    store: StoreResponse
    channel: Optional[ChannelResponse]
    delivery_service: Optional[DeliveryServiceResponse]
    departments: List[DepartmentWithSentimentResponse]
    topics: List[TopicWithSentimentResponse]
    keywords: List[KeywordWithSentimentResponse]
    comment: str
    sentiment: str
    reported_at: datetime
    created_at: datetime
    updated_at: datetime


@router.get("")
async def get_surveys(
    filter_params: FilterRequest = Depends(get_filter_params),
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=100),
    db: Session = Depends(get_db),
) -> Page[SurveyResponse]:
    filter_dict = filter_params.model_dump()
    filtered_query = build_survey_query(db.query(Survey).distinct(), filter_dict)

    # Apply ordering
    ordered_query = filtered_query.order_by(Survey.reported_at.desc())

    # Calculate total count
    total = ordered_query.count()

    # Apply pagination
    offset = (page - 1) * size
    surveys = ordered_query.offset(offset).limit(size).all()

    # Convert to response models
    survey_responses = [
        SurveyResponse.model_validate(survey.to_dict()) for survey in surveys
    ]

    # Create pagination response
    from fastapi_pagination import Params

    params = Params(page=page, size=size)

    return Page.create(items=survey_responses, total=total, params=params)


class CreateSurveyDepartmentRequest(BaseModel):
    name: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyTopicRequest(BaseModel):
    topic: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyKeywordRequest(BaseModel):
    keyword: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class CreateSurveyRequest(BaseModel):
    store_id: int
    channel: Optional[str] = None
    delivery_service: Optional[str] = None
    departments: List[CreateSurveyDepartmentRequest]
    comment: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"
    topics: List[CreateSurveyTopicRequest]
    keywords: List[CreateSurveyKeywordRequest]
    reported_at: datetime = Field(default_factory=datetime.now)


@router.post("/")
async def create_survey(
    survey_request: CreateSurveyRequest,
    db: Session = Depends(get_db),
):

    # Check if store exists
    store = db.query(Store).filter(Store.id == survey_request.store_id).first()
    if not store:
        raise HTTPException(status_code=404, detail="Store not found")
    # Check if departments exist
    departments = (
        db.query(Department)
        .filter(Department.name.in_([dept.name for dept in survey_request.departments]))
        .all()
    )
    if not departments:
        raise HTTPException(status_code=404, detail="Departments not found")
    # Check if topics exist
    topics = (
        db.query(Topic)
        .filter(Topic.topic.in_([topic.topic for topic in survey_request.topics]))
        .all()
    )
    if not topics:
        raise HTTPException(status_code=404, detail="Topics not found")
    # Initialize channel and delivery_service to None
    channel = None
    delivery_service = None

    # If channel is provided, check if channel exists
    if survey_request.channel:
        channel = (
            db.query(Channel).filter(Channel.name == survey_request.channel).first()
        )
        if not channel:
            raise HTTPException(status_code=404, detail="Channel not found")
    # If delivery service is provided, check if delivery service exists
    if survey_request.delivery_service:
        delivery_service = (
            db.query(DeliveryService)
            .filter(DeliveryService.name == survey_request.delivery_service)
            .first()
        )
        if not delivery_service:
            raise HTTPException(status_code=404, detail="Delivery service not found")
    try:
        # Create survey without departments, topics, and keywords relationships
        survey = Survey(
            store_id=survey_request.store_id,
            comment=survey_request.comment,
            sentiment=survey_request.sentiment,
            reported_at=survey_request.reported_at,
            channel_id=channel.id if channel else None,
            delivery_service_id=delivery_service.id if delivery_service else None,
        )
        db.add(survey)
        db.flush()
        # Add departments, topics, and keywords relationships
        for department in departments:
            survey_department = SurveyDepartments(
                survey_id=survey.id,
                department_id=department.id,
                sentiment=survey_request.sentiment,
            )
            db.add(survey_department)
        for request_topic in survey_request.topics:
            topic = db.query(Topic).filter(Topic.topic == request_topic.topic).first()
            survey_topic = SurveyTopics(
                survey_id=survey.id,
                topic_id=topic.id,
                sentiment=request_topic.sentiment,
            )
            db.add(survey_topic)
        for request_keyword in survey_request.keywords:
            keyword = (
                db.query(Keyword)
                .filter(Keyword.keyword == request_keyword.keyword)
                .first()
            )
            # Create keyword if it doesn't exist
            if not keyword:
                keyword = Keyword(keyword=request_keyword.keyword)
                db.add(keyword)
                db.flush()
            survey_keyword = SurveyKeywords(
                survey_id=survey.id,
                keyword_id=keyword.id,
                sentiment=request_keyword.sentiment,
            )
            db.add(survey_keyword)
        db.commit()
        db.refresh(survey)
        return survey.to_dict()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/download")
async def download_surveys(
    filter_params: FilterRequest = Depends(get_filter_params),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    filter_dict = filter_params.model_dump()
    filtered_query = build_survey_query(db.query(Survey).distinct(), filter_dict)

    def calculate_sentiment(survey):
        """Calculate sentiment based on topics using the same logic as frontend."""
        # Get topics with sentiment from the survey
        topics = [
            {"topic": survey_topic.topic.topic, "sentiment": survey_topic.sentiment}
            for survey_topic in survey.survey_topics
        ]

        # Safety check: ensure topics is an array
        if not topics or len(topics) == 0:
            return {"sentiment": "neutral", "score": 0}

        # if there are only neutral, return "neutral"
        if all(topic["sentiment"] == "neutral" for topic in topics):
            return {"sentiment": "neutral", "score": 0}

        # if there are only positive or neutral, return "positive"
        if all(topic["sentiment"] in ["positive", "neutral"] for topic in topics):
            return {"sentiment": "positive", "score": 1}

        # if there are only negative or neutral, return "negative"
        if all(topic["sentiment"] in ["negative", "neutral"] for topic in topics):
            return {"sentiment": "negative", "score": -1}

        # if there are both positive and negative, return "mix"
        positive_count = sum(1 for topic in topics if topic["sentiment"] == "positive")
        negative_count = sum(1 for topic in topics if topic["sentiment"] == "negative")
        neutral_count = sum(1 for topic in topics if topic["sentiment"] == "neutral")

        if positive_count > 0 and negative_count > 0:
            total_count = positive_count + negative_count + neutral_count
            score = round((positive_count - negative_count) / total_count, 2)
            return {"sentiment": "mix", "score": score}

        return {"sentiment": "neutral", "score": 0}

    def format_excel_value(value):
        if value is None:
            return ""
        elif isinstance(value, list):
            return "; ".join(str(item) for item in value)
        elif isinstance(value, datetime):
            return value.isoformat()
        else:
            return str(value)

    def generate_excel_file():
        # Create a workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Surveys"

        # Define headers
        headers = [
            "id",
            "store_id",
            "store_name",
            "hierarchy_level_1_name",
            "hierarchy_level_2_name",
            "hierarchy_level_3_name",
            "hierarchy_level_4_name",
            "hierarchy_level_5_name",
            "departments",
            "topics",
            "keywords",
            "comment",
            "channel",
            "delivery_service",
            "sentiment",
            "sentiment_score",
            "reported_at",
            "created_at",
            "updated_at",
        ]

        # Write headers to first row
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Stream surveys in batches using offset
        batch_size = 100
        offset = 0
        row_num = 2  # Start from row 2 (row 1 is headers)

        while True:
            # Get surveys
            surveys = (
                filtered_query.order_by(Survey.reported_at.asc())
                .offset(offset)
                .limit(batch_size)
                .all()
            )
            if not surveys:
                break
            for survey in surveys:
                csv_value = survey.to_csv()
                # Calculate sentiment based on topics
                sentiment_result = calculate_sentiment(survey)
                csv_value["sentiment"] = sentiment_result["sentiment"]
                csv_value["sentiment_score"] = sentiment_result["score"]

                # Write row values
                for col_idx, header in enumerate(headers, start=1):
                    value = format_excel_value(csv_value[header])
                    ws.cell(row=row_num, column=col_idx, value=value)
                row_num += 1
            offset += batch_size

        # Save workbook to BytesIO buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    excel_buffer = generate_excel_file()
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=surveys.xlsx"},
    )


@router.get("/{survey_id}")
async def get_survey(survey_id: int, db: Session = Depends(get_db)) -> SurveyResponse:
    filter_dict = {"ids": [survey_id]}
    query = build_survey_query(db.query(Survey).distinct(), filter_dict)
    survey = query.first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found")
    return SurveyResponse.model_validate(survey.to_dict())


class UpdateSurveyTopicRequest(BaseModel):
    topic: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyKeywordRequest(BaseModel):
    keyword: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyDepartmentRequest(BaseModel):
    name: str
    sentiment: Literal["positive", "negative", "neutral"] = "neutral"


class UpdateSurveyRequest(BaseModel):
    store_id: Optional[int] = None
    channel: Optional[str] = None
    delivery_service: Optional[str] = None
    departments: Optional[List[UpdateSurveyDepartmentRequest]] = None
    comment: Optional[str] = None
    sentiment: Optional[Literal["positive", "negative", "neutral"]] = None
    topics: Optional[List[UpdateSurveyTopicRequest]] = None
    keywords: Optional[List[UpdateSurveyKeywordRequest]] = None
    reported_at: Optional[datetime] = None


@router.put("/{survey_id}")
async def update_survey(
    survey_id: int, survey_request: UpdateSurveyRequest, db: Session = Depends(get_db)
) -> SurveyResponse:
    try:
        filter_dict = {"ids": [survey_id]}
        query = build_survey_query(db.query(Survey).distinct(), filter_dict)
        survey = query.first()
        if not survey:
            raise HTTPException(status_code=404, detail="Survey not found")
        if survey_request.store_id:
            survey.store_id = survey_request.store_id
        if survey_request.channel:
            survey.channel_id = survey_request.channel
        if survey_request.delivery_service:
            survey.delivery_service_id = survey_request.delivery_service
        if survey_request.departments:
            for department in survey_request.departments:
                survey_department = (
                    db.query(SurveyDepartments)
                    .filter(
                        SurveyDepartments.survey_id == survey.id,
                        SurveyDepartments.department_id == department.name,
                    )
                    .first()
                )
                if not survey_department:
                    survey_department = SurveyDepartments(
                        survey_id=survey.id,
                        department_id=department.name,
                        sentiment=department.sentiment,
                    )
                    db.add(survey_department)
                else:
                    survey_department.sentiment = department.sentiment
        if survey_request.comment:
            survey.comment = survey_request.comment
        if survey_request.sentiment:
            survey.sentiment = survey_request.sentiment
        if survey_request.topics:
            for topic in survey_request.topics:
                survey_topic = (
                    db.query(SurveyTopics)
                    .filter(
                        SurveyTopics.survey_id == survey.id,
                        SurveyTopics.topic_id == topic.topic,
                    )
                    .first()
                )
                if not survey_topic:
                    survey_topic = SurveyTopics(
                        survey_id=survey.id,
                        topic_id=topic.topic,
                        sentiment=topic.sentiment,
                    )
                    db.add(survey_topic)
                else:
                    survey_topic.sentiment = topic.sentiment
        if survey_request.keywords:
            for keyword in survey_request.keywords:
                survey_keyword = (
                    db.query(SurveyKeywords)
                    .filter(
                        SurveyKeywords.survey_id == survey.id,
                        SurveyKeywords.keyword_id == keyword.keyword,
                    )
                    .first()
                )
                if not survey_keyword:
                    survey_keyword = SurveyKeywords(
                        survey_id=survey.id,
                        keyword_id=keyword.keyword,
                        sentiment=keyword.sentiment,
                    )
                    db.add(survey_keyword)
                else:
                    survey_keyword.sentiment = keyword.sentiment
        if survey_request.reported_at:
            survey.reported_at = survey_request.reported_at
        db.commit()
        db.refresh(survey)
        return SurveyResponse.model_validate(survey.to_dict())
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/{survey_id}")
async def delete_survey(
    survey_id: int, db: Session = Depends(get_db)
) -> SurveyResponse:
    try:
        filter_dict = {"ids": [survey_id]}
        query = build_survey_query(db.query(Survey).distinct(), filter_dict)
        survey = query.first()
        if not survey:
            raise HTTPException(status_code=404, detail="Survey not found")
        survey.is_deleted = True
        db.commit()
        db.refresh(survey)
        return SurveyResponse.model_validate(survey.to_dict())
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


class ExtractRequest(BaseModel):
    comment: str


class ExtractedTopicResponse(BaseModel):
    text: str
    sentiment: str


class ExtractedDepartmentResponse(BaseModel):
    text: str
    sentiment: str


class ExtractedKeywordResponse(BaseModel):
    text: str
    sentiment: str


class TotalResponse(BaseModel):
    topics: list[ExtractedTopicResponse]
    departments: list[ExtractedDepartmentResponse]
    keywords: list[ExtractedKeywordResponse]
    overall_sentiment: str
    cannot_classified: bool


@router.post("/extract-total")
async def extract_total_route(
    request: ExtractRequest,
) -> tuple[TotalResponse, dict]:
    total, usage = await extract_total(request.comment)
    if total.cannot_classified:
        print("Cannot classified in AI Analysis, retrying...")
        total, usage = await extract_total_retry(request.comment)
    return total, usage
